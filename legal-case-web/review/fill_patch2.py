import openpyxl
from openpyxl.styles import Alignment

SRC = r"D:\工作用\judicial1\新套表.xlsm"
wb = openpyxl.load_workbook(SRC, keep_vba=True, data_only=False)
ws = wb["被害人筆錄"]
qs = [
    "當時的天氣和視線怎麼樣？路上有沒有號誌或標線？",
    "車禍發生後有沒有報警？警察有沒有到場？對方有留在現場嗎？",
    "事後對方有沒有跟你談賠償？有沒有和解或調解？",
]
r = ws.max_row + 1
for k, q in enumerate(qs):
    c = ws.cell(row=r + k, column=10, value=q)
    c.alignment = Alignment(vertical="center", wrap_text=True)
wb.save(SRC)
print("PATCHED col10 +3")

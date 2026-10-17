import yaml
import openpyxl
from openpyxl.styles import Alignment

SRC = r"D:\工作用\judicial1\新套表.xlsm"
REVIEW_DIR = r"D:\工作用\judicial1\legal-case-web\review"

def load(p):
    with open(p, encoding="utf-8") as f:
        d = yaml.safe_load(f)
        return d if isinstance(d, list) else [d]

master = load(REVIEW_DIR + "\\B_master_case_studies.yaml")
by_pair = {}
for c in master:
    if isinstance(c, dict) and c.get("pair_id"):
        by_pair.setdefault(c["pair_id"], []).append(c)

def qs(pid):
    out = []
    for c in by_pair.get(pid, []):
        for q in c.get("neutral_questions", []):
            q = str(q).strip()
            if q and q not in out:
                out.append(q)
    return out

FILL = {
    "涉嫌人通用": [
        "你上述年籍資料是否正確？是否為你本人？有無攜帶身分證明文件?",
        "警方所告知之罪名及權利你是否清楚？你有無刑案紀錄？",
        "你今日因何事至本所製作筆錄？",
    ],
    "過失傷害": qs("PAIR-014"),
    "撤回告訴": [
        "你是否確定要撤回告訴？是什麼原因讓你這樣決定？",
        "如果對方願意和解或接受調解，你願不願意接受？",
    ],
    "違反保護令": qs("PAIR-031"),
    "性騷擾": qs("PAIR-023"),
}

wb = openpyxl.load_workbook(SRC, keep_vba=True, data_only=False)
ws = wb["涉嫌人筆錄"]
headers = [c.value for c in ws[1]]
start_row = ws.max_row + 1
for col, questions in FILL.items():
    if col not in headers:
        print("MISS COL", col)
        continue
    j = headers.index(col) + 1
    for k, q in enumerate(questions):
        c = ws.cell(row=start_row + k, column=j, value=q)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    print(f"{col}: +{len(questions)}")
wb.save(SRC)
print("PATCH DONE")

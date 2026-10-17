import os, re, shutil
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

SRC = r"D:\工作用\judicial1\新套表.xlsm"
BAK = r"D:\工作用\judicial1\新套表_備份_20261015.xlsm"
REVIEW_DIR = r"D:\工作用\judicial1\legal-case-web\review"

import yaml
def load(p):
    with open(p, encoding="utf-8") as f:
        d = yaml.safe_load(f)
        return d if isinstance(d, list) else [d]

DATE_EXPR = 'LEFT(輸入!$C$8,3)&"年"&MID(輸入!$C$8,4,2)&"月"&MID(輸入!$C$8,6,2)&"日"'
PLACE_EXPR = '輸入!$C$7'
DATE_TRIGGERS = ["什麼時間", "什麼時候", "何時", "幾點"]
PLACE_TRIGGERS = ["什麼地點", "什麼地方", "在哪裡", "哪裡"]

def to_formula(q):
    pattern = "(" + "|".join(re.escape(t) for t in DATE_TRIGGERS + PLACE_TRIGGERS) + ")"
    tokens = re.split(pattern, q)
    if len(tokens) == 1:
        return None
    out = []
    for tok in tokens:
        if tok in DATE_TRIGGERS:
            out.append(DATE_EXPR)
        elif tok in PLACE_TRIGGERS:
            out.append(PLACE_EXPR)
        elif tok:
            out.append('"' + tok.replace('"', '""') + '"')
    return "=CONCATENATE(" + ",".join(out) + ")"

PAIR_COL = {
    "PAIR-001": "竊盜", "PAIR-002": "詐欺", "PAIR-003": "竊盜", "PAIR-004": "毀損",
    "PAIR-005": "竊盜", "PAIR-006": "侵占", "PAIR-007": "恐嚇", "PAIR-008": "詐欺",
    "PAIR-009": "竊盜", "PAIR-010": "偽造文書", "PAIR-011": "傷害", "PAIR-012": "傷害",
    "PAIR-013": "傷害", "PAIR-014": "傷害", "PAIR-015": "恐嚇", "PAIR-016": "恐嚇",
    "PAIR-017": "妨害自由", "PAIR-018": "妨害名譽", "PAIR-019": "誣告", "PAIR-020": "誣告",
    "PAIR-021": "AI_妨害性自主", "PAIR-022": "AI_妨害性自主", "PAIR-023": "性騷擾防治法",
    "PAIR-024": "AI_兒少性剝削", "PAIR-025": "妨害名譽", "PAIR-026": "妨害自由",
    "PAIR-027": "公共危險", "PAIR-028": "毒品危害防制條例", "PAIR-029": "AI_圖利賄賂",
    "PAIR-030": "跟蹤騷擾", "PAIR-031": "傷害", "PAIR-032": "妨害自由",
    "PAIR-033": "AI_少年事件", "PAIR-034": "性騷擾防治法",
    "PAIR-038": "妨害名譽", "PAIR-039": "AI_賭博", "PAIR-040": "AI_公然猥褻",
    "PAIR-041": "妨害自由", "PAIR-042": "傷害", "PAIR-043": "毀損",
    "PAIR-044": "妨害秩序", "PAIR-045": "AI_槍砲", "PAIR-046": "AI_槍砲",
    "PAIR-047": "詐欺", "PAIR-048": "詐欺", "PAIR-049": "竊盜", "PAIR-050": "詐欺",
    "PAIR-051": "傷害", "PAIR-052": "妨害電腦使用", "PAIR-053": "傷害",
    "SINGLE_THEFT": "竊盜", "SINGLE_ASSAULT": "傷害",
    "SINGLE_FORGERY": "偽造文書", "SINGLE_THREAT": "恐嚇",
}

col_qs = {}
def add(col, q):
    q = str(q).strip()
    if not q:
        return
    col_qs.setdefault(col, [])
    if q not in col_qs[col]:
        col_qs[col].append(q)

for path in [REVIEW_DIR + "\\B_master_case_studies.yaml",
             REVIEW_DIR + "\\B_single_offense_questions.yaml",
             REVIEW_DIR + "\\B9_questions_supplement.yaml"]:
    if not os.path.exists(path):
        print("SKIP", path)
        continue
    for case in load(path):
        if not isinstance(case, dict):
            continue
        col = PAIR_COL.get(case.get("pair_id", ""))
        if not col:
            continue
        for q in case.get("neutral_questions", []):
            add(col, q)

# 備份原檔（位元組完整拷貝）
shutil.copy2(SRC, BAK)
print("BACKUP ->", BAK)

wb = openpyxl.load_workbook(SRC, keep_vba=True, data_only=False)
ws = wb["涉嫌人筆錄"]
orig_headers = [c.value for c in ws[1]]
new_cols = [c for c in col_qs if c not in orig_headers]
headers = orig_headers + new_cols
for j, h in enumerate(headers, 1):
    if j > len(orig_headers):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(j)].width = 55

start_row = ws.max_row + 1
n_f = 0
for j, h in enumerate(headers, 1):
    for k, q in enumerate(col_qs.get(h, [])):
        f = to_formula(q)
        c = ws.cell(row=start_row + k, column=j, value=f if f else q)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if f:
            n_f += 1
            c.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

wb.save(SRC)
print(f"FILLED cols={len(headers)} new_cols={new_cols} formula_cells={n_f}")

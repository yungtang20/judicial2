import os, re
import yaml
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

REVIEW_DIR = r"D:\工作用\judicial1\legal-case-web\review"
SRC = r"D:\工作用\judicial1\新套表.xlsm"

DATE_EXPR = 'LEFT(輸入!$C$8,3)&"年"&MID(輸入!$C$8,4,2)&"月"&MID(輸入!$C$8,6,2)&"日"'
PLACE_EXPR = '輸入!$C$7'
DATE_TRIGGERS = ["什麼時間", "什麼時候", "何時", "幾點"]
PLACE_TRIGGERS = ["什麼地點", "什麼地方", "在哪裡", "哪裡"]

def to_formula(q):
    pattern = "(" + "|".join(re.escape(t) for t in DATE_TRIGGERS + PLACE_TRIGGERS) + ")"
    tokens = re.split(pattern, q)
    if len(tokens) == 1:
        return None
    out = []
    for tok in tokens:
        if tok in DATE_TRIGGERS:
            out.append(DATE_EXPR)
        elif tok in PLACE_TRIGGERS:
            out.append(PLACE_EXPR)
        elif tok:
            out.append('"' + tok.replace('"', '""') + '"')
    return "=CONCATENATE(" + ",".join(out) + ")"

PAIR_COL = {
    "PAIR-001": "竊盜", "PAIR-002": "詐欺", "PAIR-003": "竊盜", "PAIR-004": "毀損",
    "PAIR-005": "竊盜", "PAIR-006": "侵占", "PAIR-007": "恐嚇", "PAIR-008": "詐欺",
    "PAIR-009": "竊盜", "PAIR-010": "偽造文書", "PAIR-011": "傷害", "PAIR-012": "傷害",
    "PAIR-013": "傷害", "PAIR-014": "傷害", "PAIR-015": "恐嚇", "PAIR-016": "恐嚇",
    "PAIR-017": "妨害自由", "PAIR-018": "妨害名譽", "PAIR-019": "誣告", "PAIR-020": "誣告",
    "PAIR-021": "AI_妨害性自主", "PAIR-022": "AI_妨害性自主", "PAIR-023": "性騷擾防治法",
    "PAIR-024": "AI_兒少性剝削", "PAIR-025": "妨害名譽", "PAIR-026": "妨害自由",
    "PAIR-027": "公共危險", "PAIR-028": "毒品危害防制條例", "PAIR-029": "AI_圖利賄賂",
    "PAIR-030": "跟蹤騷擾", "PAIR-031": "傷害", "PAIR-032": "妨害自由",
    "PAIR-033": "AI_少年事件", "PAIR-034": "性騷擾防治法",
    "PAIR-038": "妨害名譽", "PAIR-039": "AI_賭博", "PAIR-040": "AI_公然猥褻",
    "PAIR-041": "妨害自由", "PAIR-042": "傷害", "PAIR-043": "毀損",
    "PAIR-044": "妨害秩序", "PAIR-045": "AI_槍砲", "PAIR-046": "AI_槍砲",
    "PAIR-047": "詐欺", "PAIR-048": "詐欺", "PAIR-049": "竊盜", "PAIR-050": "詐欺",
    "PAIR-051": "傷害", "PAIR-052": "妨害電腦使用", "PAIR-053": "傷害",
    "SINGLE_THEFT": "竊盜", "SINGLE_ASSAULT": "傷害",
    "SINGLE_FORGERY": "偽造文書", "SINGLE_THREAT": "恐嚇",
}

def load(p):
    with open(p, encoding="utf-8") as f:
        d = yaml.safe_load(f)
        return d if isinstance(d, list) else [d]

# 收集 AI 問句（依欄位分組，去重保序）
col_qs = {}
def add(col, q):
    q = str(q).strip()
    if not q:
        return
    col_qs.setdefault(col, [])
    if q not in col_qs[col]:
        col_qs[col].append(q)

for path in [REVIEW_DIR + "\\B_master_case_studies.yaml",
             REVIEW_DIR + "\\B_single_offense_questions.yaml",
             REVIEW_DIR + "\\B9_questions_supplement.yaml"]:
    if not os.path.exists(path):
        continue
    for case in load(path):
        if not isinstance(case, dict):
            continue
        pid = case.get("pair_id", "")
        col = PAIR_COL.get(pid)
        if not col:
            continue
        for q in case.get("neutral_questions", []):
            add(col, q)

# 讀原套表（計算值）
src = openpyxl.load_workbook(SRC, data_only=True)
wb = openpyxl.Workbook()

# 1. 輸入骨架頁（供公式引用）
ws_in = wb.active
ws_in.title = "輸入"
ws_in["B2"], ws_in["C2"] = "案由", ""
ws_in["B5"], ws_in["C5"] = "報案內容簡述", ""
ws_in["B7"], ws_in["C7"] = "發生地點", ""
ws_in["B8"], ws_in["C8"] = "發生日期(YYYMMDD…)", ""
ws_in["B10"], ws_in["C10"] = "報案日期", ""
ws_in["B12"], ws_in["C12"] = "證明單編號", ""
ws_in["A1"] = "請在此頁填入案件基本資料，筆錄頁問句自動帶入"
for row in ws_in.iter_rows(min_row=1, max_row=12, max_col=3):
    for c in row:
        c.alignment = Alignment(vertical="center", wrap_text=True)
ws_in.column_dimensions['B'].width = 24
ws_in.column_dimensions['C'].width = 50

# 2. 涉嫌人筆錄：原欄位 + AI 問句填入
ws0 = src["涉嫌人筆錄"]
orig_headers = [c.value for c in ws0[1]]
new_cols = [c for c in col_qs if c not in orig_headers]
headers = orig_headers + new_cols
ws = wb.create_sheet("涉嫌人筆錄")
for j, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 55

# 原問句（靜態值）
max_row = ws0.max_row
for i in range(2, max_row + 1):
    for j, h in enumerate(orig_headers, 1):
        v = ws0.cell(row=i, column=j).value
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        c = ws.cell(row=i, column=j, value=str(v))
        c.alignment = Alignment(vertical="center", wrap_text=True)

# AI 問句接續填入（公式 live）
start_row = max_row + 1
for j, h in enumerate(headers, 1):
    qs = col_qs.get(h, [])
    for k, q in enumerate(qs):
        f = to_formula(q)
        c = ws.cell(row=start_row + k, column=j, value=f if f else q)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if f:
            c.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

# 3. 被害人 / 證人頁：原樣複製
for name in ["被害人筆錄", "證人筆錄"]:
    w0 = src[name]
    w = wb.create_sheet(name)
    for row in w0.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            c = w.cell(row=cell.row, column=cell.column, value=cell.value)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    for j in range(1, w.max_column + 1):
        w.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 55

out = os.path.join(REVIEW_DIR, "新套表_AI問句填入版.xlsx")
wb.save(out)
n_ai = sum(len(v) for v in col_qs.values())
n_f = sum(1 for v in col_qs.values() for q in v if to_formula(q))
print(f"AI_Q {n_ai} FORMULA {n_f} COLS {len(headers)} -> {out}")

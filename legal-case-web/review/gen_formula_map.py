import os, re
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

REVIEW_DIR = r"D:\工作用\judicial1\legal-case-web\review"

DATE_EXPR = 'LEFT(輸入!$C$8,3)&"年"&MID(輸入!$C$8,4,2)&"月"&MID(輸入!$C$8,6,2)&"日"'
PLACE_EXPR = '輸入!$C$7'
DATE_TRIGGERS = ["什麼時間", "什麼時候", "何時", "幾點"]
PLACE_TRIGGERS = ["什麼地點", "什麼地方", "在哪裡", "哪裡"]

def to_formula(q):
    parts, used = [], []
    # 以觸發詞切分，逐段重組
    pattern = "(" + "|".join(re.escape(t) for t in DATE_TRIGGERS + PLACE_TRIGGERS) + ")"
    tokens = re.split(pattern, q)
    if len(tokens) == 1:
        return None, []  # 無槽位，不需公式化
    out = []
    for tok in tokens:
        if tok in DATE_TRIGGERS:
            out.append(DATE_EXPR)
            if "日期" not in used:
                used.append("日期(輸入!$C$8)")
        elif tok in PLACE_TRIGGERS:
            out.append(PLACE_EXPR)
            if "地點(輸入!$C$7)" not in used:
                used.append("地點(輸入!$C$7)")
        elif tok:
            out.append('"' + tok.replace('"', '""') + '"')
    return "=CONCATENATE(" + ",".join(out) + ")", used

def load(p):
    with open(p, encoding="utf-8") as f:
        d = yaml.safe_load(f)
        return d if isinstance(d, list) else [d]

rows = []
for path, is_single in [(REVIEW_DIR + "\\B_master_case_studies.yaml", False),
                         (REVIEW_DIR + "\\B_single_offense_questions.yaml", True)]:
    if not os.path.exists(path):
        continue
    for case in load(path):
        if not isinstance(case, dict):
            continue
        pid = case.get("pair_id", "")
        label = (case.get("offense", "") or pid) if is_single else pid
        for i, q in enumerate(case.get("neutral_questions", []), 1):
            f, used = to_formula(str(q))
            rows.append([label, f"Q{i}", str(q), f if f else str(q),
                         "、".join(used) if used else "無（靜態沿用）"])

wb = Workbook()
ws = wb.active
ws.title = "問句公式對照表"
headers = ["PAIR/罪名", "題號", "靜態版問句", "公式版（貼入套表儲存格）", "引用欄位"]
ws.append(headers)
fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for r in rows:
    ws.append(r)
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 8
ws.column_dimensions['C'].width = 55
ws.column_dimensions['D'].width = 55
ws.column_dimensions['E'].width = 22
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=5):
    for cell in row:
        cell.alignment = Alignment(vertical="center", wrap_text=True)

out = os.path.join(REVIEW_DIR, "問句公式對照表.xlsx")
wb.save(out)
n_formula = sum(1 for r in rows if r[4] != "無（靜態沿用）")
print(f"TOTAL {len(rows)}  FORMULA {n_formula}  STATIC {len(rows)-n_formula} -> {out}")

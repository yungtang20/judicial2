import openpyxl, shutil, re
from difflib import SequenceMatcher

SRC = r"D:\工作用\judicial1\新套表.xlsm"
BAK2 = r"D:\工作用\judicial1\新套表_備份2_20261015.xlsm"
shutil.copy2(SRC, BAK2)

SKIP_PAT = re.compile("年籍資料|三項權利|精神狀況|一人製作|刑法第57條|夜間|屬實正確|其他意見補充|自由意識|錄音|送達地址|身心障礙|低收入|原住民|親友|辯護人|法扶|提審|刑案紀錄")

def norm(s):
    s = str(s or "")
    s = re.sub(r"[0-9０-９年月日時分許第號路街段巷弄樓之0-9\-/()（）:：\.、，,？?「」\"' ]+", "", s)
    return s

wb = openpyxl.load_workbook(SRC, keep_vba=True, data_only=False)
wb_val = openpyxl.load_workbook(SRC, data_only=True)
report = []
for sheet, base in [("涉嫌人筆錄", 18), ("被害人筆錄", 27), ("證人筆錄", 18)]:
    ws, wv = wb[sheet], wb_val[sheet]
    # 原實體問句 (排除起手/收尾 markers，取顯示值)
    orig = {}
    for r in range(2, base + 1):
        for c in ws[r]:
            v = wv.cell(row=r, column=c.column).value
            t = str(v or "").strip()
            if t and not SKIP_PAT.search(t):
                orig.setdefault(c.column, []).append((r, t))
    # AI 新增列
    for r in range(base + 1, ws.max_row + 1):
        for c in ws[r]:
            t = str(c.value or "").strip()
            if not t or t.startswith("AI_"):
                continue
            if SKIP_PAT.search(t):
                continue
            hit = None
            for j, lst in orig.items():
                for (rr, ot) in lst:
                    a, b = norm(t), norm(ot)
                    if a and (a == b or (len(a) > 8 and len(b) > 8 and SequenceMatcher(None, a, b).ratio() >= 0.8)):
                        hit = (j, rr, ot)
                        break
                if hit:
                    break
            if hit:
                j, rr, ot = hit
                if c.column == j:
                    # 同欄重複：AI 格清空（保留原文）
                    c.value = None
                    report.append(f"[{sheet}] 同欄重複，已清空AI格 R{r}C{c.column}：{t[:30]}… ≒原文R{rr}")
                else:
                    # 跨欄相近：保留兩者（不同罪名脈絡），僅記錄
                    report.append(f"[{sheet}] 跨欄相近保留 R{r}C{c.column} vs 原R{rr}C{j}：{t[:30]}…")
print(f"DUP_HITS {len(report)}")
open(r"C:\Users\yungtang\AppData\Local\Temp\opencode\dup.txt", "w", encoding="utf-8").write("\n".join(report) if report else "無重複")
wb.save(SRC)
print("SAVED", BAK2)

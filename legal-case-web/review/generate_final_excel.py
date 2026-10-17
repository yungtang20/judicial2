import os
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

REVIEW_DIR = r"D:\工作用\judicial1\legal-case-web\review"

# 1. 定義通用問句 (從 extracted_excel_questions.md 提取之標準模板)
GENERIC_START = [
    "你上述年籍資料是否正確？是否為你本人？有無攜帶身分證明文件?",
    "警方所告知之罪名及權利你是否清楚？你有無刑案紀錄？",
    "你今日因何事至本所製作筆錄？"
]
GENERIC_END = [
    "以上筆錄所述登載是否屬實正確？有無其他意見補充？",
    "警方於偵訊時你所陳述之事實是否為你自由意識下所陳述？警方有無以強暴脅迫等不正方法對你取供？"
]

# 2. 讀取 YAML 檔案
def load_yaml(path):
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            data = yaml.safe_load(f)
            return data if isinstance(data, list) else [data]
    return []

master_data = load_yaml(os.path.join(REVIEW_DIR, "B_master_case_studies.yaml"))
single_data = load_yaml(os.path.join(REVIEW_DIR, "B_single_offense_questions.yaml"))

# 3. 建立 Excel 工作簿
wb = Workbook()
ws = wb.active
ws.title = "罪章筆錄問句知識庫"

# 設定欄位標題
headers = ["第一欄：罪章 / PAIR 分類", "候選罪名 A", "候選罪名 B", "核心區辨焦點 / 法律目標", "問句階段", "問句內容 (底層事實/通用)"]
ws.append(headers)

# 設定標題列樣式
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 4. 處理資料並寫入 Excel
processed_pairs = set() # 用於去重

def process_case(case, is_single=False):
    if not isinstance(case, dict): return

    if is_single:
        pair_id = case.get("pair_id", "Unknown")
        offense = case.get("offense", "Unknown")
        category = f"單一罪名: {offense}"
        cand_a = offense
        cand_b = "N/A"
        focus = case.get("legal_issue", "")
        questions = case.get("neutral_questions", [])
    else:
        pair_id = case.get("pair_id", "Unknown")
        if pair_id in processed_pairs: return # 去重
        processed_pairs.add(pair_id)

        category = f"{pair_id}"
        legal_analysis = case.get("legal_analysis", {})
        cand_a = legal_analysis.get("candidate_a", {}).get("name", "")
        cand_b = legal_analysis.get("candidate_b", {}).get("name", "")
        focus = case.get("legal_issue", "")
        if not focus:
            df = case.get("distinction_facts", "")
            if isinstance(df, list):
                df = "\n".join(df)
            focus = str(df)[:150] + ("..." if len(str(df)) > 150 else "")
        questions = case.get("neutral_questions", [])

    # 寫入通用起手式
    for q in GENERIC_START:
        ws.append([category, cand_a, cand_b, focus, "通用起手式 (程序與權利)", q])

    # 寫入實體區辨問句
    for q in questions:
        ws.append([category, cand_a, cand_b, focus, "實體區辨問句 (底層事實)", q])

    # 寫入通用收尾式
    for q in GENERIC_END:
        ws.append([category, cand_a, cand_b, focus, "通用收尾式 (意願與送達)", q])

# 處理混淆組合
for case in master_data:
    process_case(case, is_single=False)

# 處理單一罪名
for case in single_data:
    process_case(case, is_single=True)

# 5. 調整欄寬與格式
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 15
ws.column_dimensions['D'].width = 40
ws.column_dimensions['E'].width = 25
ws.column_dimensions['F'].width = 60

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=6):
    for cell in row:
        cell.alignment = Alignment(vertical="center", wrap_text=True)

# 6. 儲存檔案
output_excel = os.path.join(REVIEW_DIR, "最終版_罪章筆錄問句知識庫.xlsx")
wb.save(output_excel)

print(f"=== 轉換完成 ===")
print(f"已成功生成: {output_excel}")
print(f"共處理 {len(processed_pairs)} 組 PAIR 及 {len(single_data)} 個單一罪名。")

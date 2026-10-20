# -*- coding: utf-8 -*-
"""新套表.xlsm 未入庫抽題差異檢查（review-only 工具）。

背景：`data/source_manifest.json` 標示工作簿於 2026-10-15 增補
506 格題目，`data/xlsm_questions.json` 未重抽，新增題目尚未人工
覆核、不得標為 COMPANY_VERBATIM、未進入 runtime。

本工具不修改任何 runtime 資料：它只比對三個筆錄 sheet
（被害人筆錄／涉嫌人筆錄／證人筆錄）中尚未出現在既存抽取快取的
文字格，產出人工覆核草稿（markdown），供後續逐題核定。

用法：
    py -3 -X utf8 tools/xlsm_review_diff.py

輸出：review/xlsm_review_draft.md（review-only，未經核定不得加載）
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
CONFIG = {
    "workbook_path": REPO_ROOT.parent / "新套表.xlsm",
    "cache_path": REPO_ROOT / "data" / "xlsm_questions.json",
    "output_path": REPO_ROOT / "review" / "xlsm_review_draft.md",
}
SHEET_TO_ROLE = {
    "被害人筆錄": "victim",
    "涉嫌人筆錄": "suspect",
    "證人筆錄": "witness",
}
DRAFT_HEADER_NOTE = (
    "未入 runtime：新增格在人工覆核核定前不得標為 COMPANY_VERBATIM，"
    "不得加入 data/xlsm_questions.json，亦不得出現在 build_questions 輸出。"
    "AI_* 欄（AI_少年事件、AI_賭博、AI_槍砲等）為 AI 輔助生成題，"
    "核定後標為 AI_GENERATED，與公司既有套表題（COMPANY_VERBATIM）分開。"
)
REVIEW_STATUS = "HUMAN_REVIEW_REQUIRED"


def load_cache_texts(cache_path: Path) -> set[str]:
    import json

    cache = json.loads(cache_path.read_text(encoding="utf-8"))
    texts: set[str] = set()
    for role_items in cache.values():
        if not isinstance(role_items, list):
            continue
        for item in role_items:
            if isinstance(item, dict) and isinstance(item.get("text"), str):
                texts.add(item["text"].strip())
    return texts


def _is_formula(value: str) -> bool:
    return value.lstrip().startswith("=")


def collect_pending(workbook_path: Path, cache_texts: set[str]) -> list[dict[str, Any]]:
    """收集三個筆錄 sheet 中不在快取內的文字格（含所在罪章欄名）。"""
    wb = openpyxl.load_workbook(workbook_path, data_only=False, read_only=True)
    pending: list[dict[str, Any]] = []
    try:
        for sheet_name, role in SHEET_TO_ROLE.items():
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            headers: dict[str, str] = {}
            for row in ws.iter_rows(min_row=1, max_row=1):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.strip():
                        headers[cell.column_letter] = cell.value.strip()
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    text = cell.value.strip()
                    if not text or _is_formula(text):
                        continue
                    if text in cache_texts:
                        continue
                    pending.append(
                        {
                            "sheet": sheet_name,
                            "role": role,
                            "row": cell.row,
                            "col": cell.column_letter,
                            "offense_header": headers.get(cell.column_letter, "未標示罪章"),
                            "text": text,
                        }
                    )
    finally:
        wb.close()
    pending.sort(key=lambda item: (item["role"], item["col"], item["row"]))
    return pending


def render_draft(pending: list[dict[str, Any]]) -> str:
    lines: list[str] = [
        "# 新套表.xlsm 未入庫抽題覆核草稿",
        "",
        f"- 覆核狀態：**{REVIEW_STATUS}**",
        f"- 共 {len(pending)} 格尚未出現在 data/xlsm_questions.json（含原本就未抽取的通用格）。",
        f"- {DRAFT_HEADER_NOTE}",
        "- 差異工具：`tools/xlsm_review_diff.py`（可重跑比對）。",
        "",
    ]
    current_header = None
    for item in pending:
        key = (item["role"], item["col"], item["offense_header"])
        if key != current_header:
            current_header = key
            is_ai = item["offense_header"].startswith("AI_")
            tag = "（AI 輔助生成）" if is_ai else ""
            lines.append(
                f"## {item['sheet']}（{item['role']}）第 {item['col']} 欄：{item['offense_header']}{tag}"
            )
        lines.append(f"- 第 {item['row']} 列｜原文：「{item['text']}」")
    if not pending:
        lines.append("（無待覆核格）")
    lines.append("")
    return "\n".join(lines)


def main() -> int:
    cache_texts = load_cache_texts(CONFIG["cache_path"])
    pending = collect_pending(CONFIG["workbook_path"], cache_texts)
    markdown = render_draft(pending)
    output: Path = CONFIG["output_path"]
    output.write_text(markdown, encoding="utf-8")
    print(f"pending_cells={len(pending)} -> {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

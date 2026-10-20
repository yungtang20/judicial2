"""xlsm_review_diff 工具單元測試：只測純函數，不連真實工作簿。

抽取差異產出必須維持 review-only 契約：
- 只回報未入快取的文字格；
- 跳過含公式（= 開頭）的格；
- 草稿明確標示 HUMAN_REVIEW_REQUIRED 且說明不得進入 runtime。
"""

from __future__ import annotations

import openpyxl
import pytest

from tools.xlsm_review_diff import (
    REVIEW_STATUS,
    SHEET_TO_ROLE,
    collect_pending,
    load_cache_texts,
    render_draft,
)


@pytest.fixture()
def mini_workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "被害人筆錄"
    ws["A1"] = "被害人通用"
    ws["B1"] = "竊盜"
    ws["A2"] = "你今（  ）日因何事接受警方製作筆錄？"
    ws["B2"] = "新分店防損題：收銀機短溢差如何確認？"
    ws["B3"] = "=CONCATENATE(\"今日\",輸入!A1)"
    path = tmp_path / "mini.xlsm"
    wb.save(path)
    return path


@pytest.fixture()
def mini_cache(tmp_path):
    cache = {
        "victim": [
            {"row": 2, "col": "A", "offense": "被害人通用", "text": "你今（  ）日因何事接受警方製作筆錄？"}
        ],
        "suspect": [],
        "witness": [],
    }
    path = tmp_path / "cache.json"
    path.write_text(
        import_json_cache(cache), encoding="utf-8"
    )
    return path


def import_json_cache(cache: dict) -> str:
    import json

    return json.dumps(cache, ensure_ascii=False)


def test_collect_pending_skips_cached_and_formula_cells(mini_workbook, mini_cache):
    texts = load_cache_texts(mini_cache)
    pending = collect_pending(mini_workbook, texts)
    texts_found = [item["text"] for item in pending]
    # 已入快取的 A2 與公式格 B3 都不得视为待覆核
    assert texts_found == ["新分店防損題：收銀機短溢差如何確認？"]
    item = pending[0]
    assert item["role"] == "victim"
    assert item["offense_header"] == "竊盜"
    assert item["row"] == 2
    assert item["col"] == "B"


def test_render_draft_marks_review_only_contract(mini_workbook, mini_cache):
    texts = load_cache_texts(mini_cache)
    pending = collect_pending(mini_workbook, texts)
    draft = render_draft(pending)
    assert REVIEW_STATUS in draft
    assert "COMPANY_VERBATIM" in draft
    assert "data/xlsm_questions.json" in draft
    assert "tools/xlsm_review_diff.py" in draft
    # 不得冒稱任何已核定狀態
    assert "VERIFIED" not in draft.replace("HUMAN_REVIEW_REQUIRED", "")


def test_render_draft_marks_ai_columns_with_tag(tmp_path):
    """AI_ 開頭欄名在草稿中必須標為「AI 輔助生成」，不得混入 COMPANY_VERBATIM。"""
    pending = [
        {
            "sheet": "涉嫌人筆錄", "role": "suspect", "row": 20, "col": "AA",
            "offense_header": "AI_少年事件", "text": "行為人行為時是否滿18歲？",
        },
        {
            "sheet": "涉嫌人筆錄", "role": "suspect", "row": 20, "col": "B",
            "offense_header": "過失傷害", "text": "你把提款卡交給對方前有沒有想過？",
        },
    ]
    draft = render_draft(pending)
    assert "AI_少年事件（AI 輔助生成）" in draft
    assert "AI_GENERATED" in draft
    # 非 AI 欄不得被標為 AI
    assert "過失傷害（AI 輔助生成）" not in draft


def test_sheet_role_map_covers_runtime_sheets():
    assert set(SHEET_TO_ROLE) == {"被害人筆錄", "涉嫌人筆錄", "證人筆錄"}

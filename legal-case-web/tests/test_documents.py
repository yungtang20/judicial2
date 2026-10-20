from __future__ import annotations

from pathlib import Path

import io
import zipfile

from docx import Document

from app.documents import brief_text, checked_interview_docx, interview_docx, _format_tw_time
from tests.helpers import completed_case


def test_format_tw_time_blocks_unparseable_free_text_in_formal_documents():
    """正式文件時間門檻：ISO 與民國樣式原樣／轉換呈現；自由文字回退
    UNKNOWN（例：「昨天下午(忘記幾點)」不得直接印入陳報單）。"""
    assert _format_tw_time("2026-10-15T10:30") == "115年10月15日10時30分"
    assert _format_tw_time("115年10月15日10時許") == "115年10月15日10時許"
    assert _format_tw_time("民國115年10月15日") == "民國115年10月15日"
    assert _format_tw_time("UNKNOWN") == "UNKNOWN"
    assert _format_tw_time("昨天下午(忘記幾點)") == "UNKNOWN"
    assert _format_tw_time("garbage") == "UNKNOWN"


def test_brief_renders_unknown_for_free_text_event_time():
    """探針 I 系列回歸：free-text event_time 曾直達陳報單「二、發生時間」。
    現在文件層必須回退 UNKNOWN，交由人工補填，不以臆測時間入正式文件。"""
    from tests.helpers import intake, BANK
    from app.workflow import create_case, apply_action

    info = intake()
    info["event_time"] = "昨天下午(忘記幾點)"
    case = create_case(info, BANK)
    case = apply_action(case, "auto_prepare", {}, BANK)
    case = apply_action(case, "save_profile", {
        "profile": {"name": "測試受詢問人", "national_id": "A123456789", "sex": "男", "education": "大學"},
        "session": {"start_time": "2026-10-11T14:00", "end_time": "2026-10-11T15:00", "location": "測試所"},
    }, BANK)
    answers = {q["id"]: "測試回答" for q in case["questions"]}
    case = apply_action(case, "save_answers", {"answers": answers}, BANK)
    case = apply_action(case, "approve_gate2", {"confirmed": True}, BANK)
    case = apply_action(case, "save_case_numbers", {
        "case_numbers": {"proof": "Z123456789", "women_child": "", "case_165": ""},
    }, BANK)
    text = brief_text(case, "victim")
    assert "二、發生時間：UNKNOWN" in text
    assert "昨天下午" not in text


def test_docx_is_openable_and_contains_all_confirmed_questions():
    case = completed_case()
    data = interview_docx(case, "victim")
    document = Document(io.BytesIO(data))
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n" + table_text
    for question in case["interviews"]["victim:default"]["questions"]:
        assert question["text"] in text
        assert question["answer"] in text
    # 公司範本結構：標題、案由、受詢問人、結尾簽名區
    assert "偵　　詢　　筆　　錄" in text
    assert "案由" in text
    assert "受詢問人" in text
    assert "並經受詢問人親閱無訛後始簽名捺印" in text
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "MERGEFIELD" not in xml
    assert "{{" not in xml


def test_docx_contains_company_required_questions_verbatim():
    case = completed_case()
    document = Document(io.BytesIO(interview_docx(case, "victim")))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    from app.company_templates import COMPANY_QUESTIONS

    # 佔位符（{today}）渲染為全形空白（　）；比對時先移除佔位符與空白
    for question in COMPANY_QUESTIONS["victim"]:
        expected = question.replace("{today}", "").replace("{now_time}", "")
        assert expected.replace("　", "") in text.replace("　", "")


def test_suspect_docx_contains_rights_notice():
    case = completed_case("suspect")
    document = Document(io.BytesIO(interview_docx(case, "suspect")))
    text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "得保持沉默，無須違背自己之意思而為陳述" in text
    assert "提審" not in text  # 有被害人涉嫌人範本無提審條


def test_transfer_docx_contains_habeas_notice():
    case = completed_case("transfer")
    document = Document(io.BytesIO(interview_docx(case, "transfer")))
    text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "提審法第1條" in text
    text_all = text + "\n" + "\n".join(p.text for p in document.paragraphs)
    assert "警方在逮捕你時，有無使你受傷" in text_all


def test_brief_has_conditional_case_numbers_and_short_summary():
    case = completed_case(offense="fraud")
    text = brief_text(case, "victim")
    assert "證明單案號" in text
    assert "165案號" in text
    assert "婦幼案號" not in text
    assert len(case["brief_summary"]) <= 500


def test_real_local_renderer_verifies_page_limit():
    case = completed_case()
    data, pages = checked_interview_docx(case, "victim")
    assert data.startswith(b"PK")
    assert 1 <= pages <= 8


def test_interview_completeness_validates_case_and_content():
    from app.document_validation import validate_interview_completeness
    from tests.helpers import BANK, intake
    from app.workflow import create_case
    case = completed_case()
    assert validate_interview_completeness(case, "victim") == []
    # 未完成 Gate 之案件應報缺失
    fresh = create_case(intake(), BANK)
    problems = validate_interview_completeness(fresh, "victim")
    assert problems and "尚未完成" in problems[0]


def test_brief_completeness_separate_from_interview():
    from app.document_validation import validate_brief_completeness
    case = completed_case(offense="fraud")
    problems = validate_brief_completeness(case, "victim")
    assert problems == []
    # 逼出長度上限
    from app.rules import concise_summary
    case["brief_summary"] = concise_summary("字" * 500)
    assert len(case["brief_summary"]) == 500


def test_docx_structure_detects_invalid_bytes():
    from app.document_validation import validate_docx_structure
    assert validate_docx_structure(b"not a zip") != []
    assert validate_docx_structure(b"PK\x03\x04-empty") != []


def test_document_output_dispatch():
    from app.document_validation import validate_document_output
    case = completed_case()
    assert validate_document_output(case, "victim", "interview") == []
    assert validate_document_output(case, "victim", "brief") == []

def test_document_validation_blocks_when_required_numbers_missing():
    from app.document_validation import validate_interview_completeness
    case = completed_case()
    # 移除證明單案號 → 應阻擋
    case["case_numbers"]["proof"] = ""
    problems = validate_interview_completeness(case, "victim")
    assert any("缺少必要案號" in p for p in problems)


def test_document_validation_blocks_when_answer_missing():
    from app.document_validation import validate_interview_completeness
    case = completed_case()
    interview = case["interviews"]["victim:default"]
    interview["questions"][0]["answer"] = ""
    problems = validate_interview_completeness(case, "victim")
    assert any("問句未作答" in p for p in problems)


def test_document_validation_blocks_suspect_without_session():
    from app.document_validation import validate_interview_completeness
    case = completed_case("suspect")
    case["interview_sessions"]["suspect:default"]["start_time"] = ""
    problems = validate_interview_completeness(case, "suspect")
    assert any("詢問起訖" in p for p in problems)

def test_interview_header_does_not_contain_candidate_marker():
    """F37：offense name 含「（候選）」會寫進筆錄告知 header（你涉嫌　加重竊盜（候選）　案）。
    候選狀態由 status/notice 傳達，不得進入罪名文字。"""
    from app.company_templates import notice_lines
    import json
    bank = json.loads((Path(__file__).resolve().parents[1] / "data" / "question_bank.json").read_text(encoding="utf-8"))
    for offense in bank["offenses"]:
        assert "（候選）" not in offense["name"], f"{offense['id']} 罪名含候選標記"
    # 渲染驗證
    header = notice_lines("suspect", "加重竊盜")[0]
    assert "（候選）" not in header

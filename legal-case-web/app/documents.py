from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pypdf import PdfReader

from .company_templates import (
    FOOTER_LINES,
    IDENTITY_FIELDS,
    INCOME_LEVELS,
    notice_lines,
)
from .config import MAX_DOCX_PAGES, MAX_SUMMARY_LENGTH
from .roc_time import format_roc_date, format_roc_time


class DocumentError(RuntimeError):
    pass


ROLE_LABELS = {
    "victim": "被害人",
    "suspect": "涉嫌人",
    "transfer": "涉嫌人（移送）",
}
ROLE_TITLES = {
    "victim": "偵　　詢　　筆　　錄　　第　1　次",
    "suspect": "偵　　詢　　筆　　錄　　第　1　次",
    "transfer": "偵　　詢　　筆　　錄　　第　1　次",
}


def _set_font(run: Any, size: int = 12, bold: bool = False) -> None:
    run.font.name = "標楷體"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
    run.font.size = Pt(size)
    run.bold = bold


def _paragraph(document: Document, text: str = "", *, bold: bool = False, center: bool = False) -> Any:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.15
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(paragraph.add_run(text), bold=bold)
    return paragraph


def _cell_text(cell: Any, text: str, *, bold: bool = False, size: int = 10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    _set_font(paragraph.add_run(text), size=size, bold=bold)


def _set_cell_width(cell: Any, cm: float) -> None:
    cell.width = Cm(cm)


def _base_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    style = document.styles["Normal"]
    style.font.name = "標楷體"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
    style.font.size = Pt(12)
    return document


def _interview(case: dict[str, Any], role: str) -> dict[str, Any]:
    interview = case.get("interviews", {}).get(f"{role}:default")
    if not interview:
        raise DocumentError("此身分尚未完成並核准問答")
    return interview


def _offense_names(interview: dict[str, Any]) -> str:
    names = []
    for question in interview["questions"]:
        name = question.get("offense_name")
        if name and name not in names:
            names.append(name)
    return "、".join(names) or "待確認"


def _format_tw_time(value: str) -> str:
    """正式文件時間門檻：ISO 可解析 → 民國格式；本身即民國樣式
    （Gate 1 摘要擷取）→ 原樣；其餘自由文字（例：「昨天下午(忘記幾點)」）
    不得直接印入正式文件，回退 UNKNOWN 由人工補填。"""
    formatted = format_roc_time(value)
    text = (value or "").strip()
    if formatted == text and text and text != "UNKNOWN" and not _TIME_PASSTHROUGH.match(text):
        return "UNKNOWN"
    return formatted


# 民國樣式時間：115年10月15日10時許／民國115年10月15日／115年10月10日10時30分
_TIME_PASSTHROUGH = re.compile(
    r"^(?:民國前)?(?:民國)?\d{1,3}年\d{1,2}月\d{1,3}日(?:\d{1,2}時(?:\d{1,2}分?)?)?許?$"
)


def _format_roc_date(value: str) -> str:
    return format_roc_date(value)


def _interview_header_table(document: Document, case: dict[str, Any], role: str, offense_name: str) -> None:
    intake = case["intake"]
    profile = case.get("profiles", {}).get(f"{role}:default", {})
    session = case.get("interview_sessions", {}).get(f"{role}:default", {})

    title = document.add_table(rows=1, cols=1)
    title.style = "Table Grid"
    title.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_text(title.rows[0].cells[0], ROLE_TITLES[role], bold=True, size=14)

    table = document.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    start = session.get("start_time") or intake["event_time"]
    end = session.get("end_time") or ""
    row = table.add_row()
    _cell_text(row.cells[0], "詢問", bold=True)
    _cell_text(row.cells[1], "時間", bold=True)
    _cell_text(row.cells[2], f"自{_format_tw_time(start)}起")
    _cell_text(row.cells[3], f"至{_format_tw_time(end)}止" if end else "至　年　月　日　時　分止")
    row = table.add_row()
    row.cells[0].merge(row.cells[0])
    _cell_text(row.cells[0], "詢問", bold=True)
    _cell_text(row.cells[1], "地點", bold=True)
    merged = row.cells[2].merge(row.cells[3])
    _cell_text(merged, intake.get("interview_location") or intake["location"])
    row = table.add_row()
    _cell_text(row.cells[0], "案由", bold=True)
    merged = row.cells[1].merge(row.cells[1])
    merged = row.cells[1]
    _cell_text(row.cells[1], f"{offense_name}({ROLE_LABELS[role]})")
    row.cells[2].text = ""
    row.cells[3].text = ""
    # 受詢問人欄
    for field in IDENTITY_FIELDS:
        row = table.add_row()
        _cell_text(row.cells[0], "受詢問人", bold=True)
        _cell_text(row.cells[1], field["label"], bold=True)
        merged = row.cells[2].merge(row.cells[3])
        _cell_text(merged, str(profile.get(field["key"], "") or ""))
    # 家庭經濟狀況
    row = table.add_row()
    _cell_text(row.cells[0], "受詢問人", bold=True)
    _cell_text(row.cells[1], "家庭經濟狀況", bold=True)
    merged = row.cells[2].merge(row.cells[3])
    _cell_text(merged, "　".join(INCOME_LEVELS) + "（承辦人勾選）")
    # 應告知事項
    if role in ("suspect", "transfer"):
        row = table.add_row()
        first = row.cells[0]
        first.text = ""
        paragraph = first.paragraphs[0]
        for char in "應告知事項":
            _set_font(paragraph.add_run(char), bold=True, size=10)
            paragraph.add_run().add_break()
        merged = row.cells[1].merge(row.cells[2]).merge(row.cells[3])
        lines = notice_lines(role, offense_name)
        merged.text = ""
        for index, line in enumerate(lines):
            paragraph = merged.paragraphs[0] if index == 0 else merged.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            _set_font(paragraph.add_run(line), size=10)


def _append_qa(document: Document, interview: dict[str, Any]) -> None:
    for index, question in enumerate(interview["questions"], start=1):
        _paragraph(document, f"問 {question['text']}")
        _paragraph(document, f"答 {question['answer'] or '　'}")


def interview_docx(case: dict[str, Any], role: str) -> bytes:
    interview = _interview(case, role)
    offense_name = _offense_names(interview)
    document = _base_document()
    _interview_header_table(document, case, role, offense_name)
    _paragraph(document)
    _append_qa(document, interview)
    for line in FOOTER_LINES:
        _paragraph(document, line)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def brief_text(case: dict[str, Any], role: str) -> str:
    intake = case["intake"]
    interview = _interview(case, role)
    names = _offense_names(interview)
    numbers = case["case_numbers"]
    case_number_parts = [f"證明單案號：{numbers['proof']}"]
    if numbers.get("women_child"):
        case_number_parts.append(f"婦幼案號：{numbers['women_child']}")
    if numbers.get("case_165"):
        case_number_parts.append(f"165案號：{numbers['case_165']}")
    report = f"依規定陳報（{'、'.join(case_number_parts)}）"
    summary = case["brief_summary"]
    if len(summary) > MAX_SUMMARY_LENGTH:
        raise DocumentError(f"條陳案情摘要超過 {MAX_SUMMARY_LENGTH} 字")
    if role == "victim":
        return (
            f"{intake['station']}報告\n"
            f"一、受理案由：{names}(受理)\n"
            f"二、發生時間：{_format_tw_time(intake['event_time'])}\n"
            f"三、發生地點：{intake['location']}\n"
            f"四、報案時間：{_format_tw_time(intake.get('report_time')) if intake.get('report_time') else 'UNKNOWN'}\n"
            f"五、案情摘要：{summary}\n\n"
            f"六、後續處置：\n（一）{report}\n"
            f"（二）本案管轄：{intake.get('jurisdiction') or 'UNKNOWN'}\n"
            f"（三）處理員警：警員 {intake['officer']}。"
        )
    if role == "suspect":
        return (
            f"{intake['station']}報告\n"
            f"一、受理案由：{names}(破獲)\n"
            f"二、發生時間：{_format_tw_time(intake['event_time'])}\n"
            f"三、發生地點：{intake['location']}\n"
            f"四、報案時間：{_format_tw_time(intake.get('report_time')) if intake.get('report_time') else 'UNKNOWN'}\n"
            f"五、破獲時間：{_format_tw_time(intake.get('arrest_time')) if intake.get('arrest_time') else 'UNKNOWN'}\n"
            f"六、案情摘要：{summary}\n\n"
            f"七、後續處置：\n（一）{report}\n"
            f"（二）後續偵辦情形：{intake.get('disposition') or 'UNKNOWN'}\n"
            f"（三）處理員警：警員 {intake['officer']}。"
        )
    return (
        f"{intake['station']}報告\n"
        f"一、查獲案由：{names}(破獲)\n"
        f"二、查獲時間：{_format_tw_time(intake['event_time'])}\n"
        f"三、查獲地點：{intake['location']}\n"
        f"四、案情摘要：{summary}\n\n"
        f"五、處理員警：警員 {intake['officer']}。"
    )


def brief_docx(case: dict[str, Any], role: str) -> bytes:
    document = _base_document()
    for line in brief_text(case, role).splitlines():
        _paragraph(document, line)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _soffice_path() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    candidate = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
    return str(candidate) if candidate.exists() else None


def verified_page_count(docx_bytes: bytes) -> int:
    executable = _soffice_path()
    if not executable:
        raise DocumentError("找不到本機 LibreOffice，無法驗證 8 頁限制")
    with tempfile.TemporaryDirectory(prefix="judicial-docx-") as directory:
        root = Path(directory)
        source = root / "document.docx"
        source.write_bytes(docx_bytes)
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "pdf", "--outdir", str(root), str(source)],
            capture_output=True,
            text=True,
            timeout=45,
            check=False,
        )
        pdf = root / "document.pdf"
        if result.returncode != 0 or not pdf.exists():
            raise DocumentError("LibreOffice 無法完成頁數驗證")
        return len(PdfReader(str(pdf)).pages)


def checked_interview_docx(case: dict[str, Any], role: str) -> tuple[bytes, int]:
    data = interview_docx(case, role)
    pages = verified_page_count(data)
    if pages > MAX_DOCX_PAGES:
        raise DocumentError(f"文件實際排版為 {pages} 頁，超過 {MAX_DOCX_PAGES} 頁上限")
    return data, pages
